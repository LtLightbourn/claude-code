"""Convert the drafter's markdown output to a .docx via python-docx."""

import re

from docx import Document
from docx.shared import Pt

_INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def markdown_to_docx(markdown_text, output_path):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            document.add_heading(_strip_inline(line[4:]), level=3)
        elif line.startswith("## "):
            document.add_heading(_strip_inline(line[3:]), level=2)
        elif line.startswith("# "):
            document.add_heading(_strip_inline(line[2:]), level=1)
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            _add_runs(paragraph, line[2:])
        elif re.match(r"^\s*[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            _add_runs(paragraph, re.sub(r"^\s*\d+\.\s+", "", line))
        else:
            paragraph = document.add_paragraph()
            _add_runs(paragraph, line)

    document.save(output_path)


def _add_runs(paragraph, text):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _strip_inline(text):
    return text.replace("**", "").replace("*", "").strip()
